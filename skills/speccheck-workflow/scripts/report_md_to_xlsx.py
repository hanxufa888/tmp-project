#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""report_md_to_xlsx.py

Step4(結果集約・レポート出力) 用スクリプト。

Step4 で生成したチェック結果レポート(Markdown, 例: `specification_check/output/diff_summary_report.md`)
内の結果テーブルを取り込み、Excel(.xlsx) へ変換する。

対象とするテーブル:
- 既定では見出し「## 結果出力」直下の Markdown テーブルを変換する(`--heading` で変更可)。
- 見出しが見つからない場合は、ファイル内で最初に出現する Markdown テーブルを変換する。

Markdown テーブルの想定:
    | 列1 | 列2 | ... |
    | --- | --- | ... |   ← 区切り行(セルが - : 空白のみ)
    | 値1 | 値2 | ... |
- 1 つの結果(セル群)は Markdown テーブルの 1 行で記載する(物理的な複数行に分割しない)。
- セル内のエスケープされたパイプ(`\\|`) はリテラルのパイプとして扱う。
- セル内の改行は `<br>` / `<br/>` / `<br />`(大文字小文字不問) で表現する。
  これらは Excel セル内の実改行(`\n`)へ変換して出力する。
- セル値の要約・改変は行わず、原文のまま出力する(前後空白のみ除去、上記 `\\|`・`<br>` の変換を除く)。

使い方（プロジェクトの .claude を起点に実行。cd は不要。パスは .claude からの相対）:
    python skills/speccheck-workflow/scripts/report_md_to_xlsx.py --in "specification_check/output/diff_summary_report.md" --out "specification_check/output/diff_summary_report.xlsx"

依存: openpyxl（Excel(.xlsx) 書き込み用）。
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.stderr.write(
        "[error] openpyxl が見つかりません。`pip install openpyxl` を実行してください。\n"
    )
    raise SystemExit(2)

# 既定で変換対象とするテーブルの見出しテキスト(先頭の # は問わない)。
DEFAULT_HEADING = "結果出力"

# 出力Excelの体裁。
MAX_COL_WIDTH = 60  # 列幅の上限(文字数換算)
MIN_COL_WIDTH = 12  # 列幅の下限(文字数換算)


def split_table_row(line: str) -> list[str]:
    """Markdown テーブル行を，エスケープされていないパイプで分割する。

    先頭/末尾のパイプに由来する空要素は除去する。`\\|` はリテラル `|` へ復元する。
    セル内の `<br>` / `<br/>` / `<br />`(大文字小文字不問) は実改行(`\\n`)へ変換する。
    """
    # エスケープされたパイプを一時的に退避してから分割する
    placeholder = "\x00"
    tmp = line.replace("\\|", placeholder)
    parts = tmp.split("|")
    # 先頭/末尾のパイプによる空要素を除去
    if parts and parts[0].strip() == "":
        parts = parts[1:]
    if parts and parts[-1].strip() == "":
        parts = parts[:-1]
    # `\\|` をリテラル `|` へ復元し，`<br>` 系タグを実改行へ変換する
    return [
        re.sub(r"(?i)<br\s*/?>", "\n", p.replace(placeholder, "|")).strip()
        for p in parts
    ]


def is_table_line(line: str) -> bool:
    return line.lstrip().startswith("|")


def is_separator_row(cells: list[str]) -> bool:
    """区切り行(セルが - : 空白のみ)かどうかを判定する。"""
    if not cells:
        return False
    return all(
        re.fullmatch(r":?-+:?", c.replace(" ", "")) is not None and c.strip() != ""
        for c in cells
    )


def find_heading_line(lines: list[str], heading: str) -> int:
    """見出し行(`#...` に heading を含む)のインデックスを返す。見つからなければ -1。"""
    norm = heading.strip()
    for i, line in enumerate(lines):
        if line.lstrip().startswith("#") and norm in line:
            return i
    return -1


def extract_table(lines: list[str], start: int) -> tuple[list[str], list[list[str]]]:
    """`start` 以降で最初に現れる Markdown テーブルを抽出する。

    戻り値: (ヘッダーセル, データ行のリスト)。テーブルが無ければ ([], [])。
    """
    # テーブルの開始行まで進める
    i = start
    n = len(lines)
    while i < n and not is_table_line(lines[i]):
        i += 1
    if i >= n:
        return [], []

    # 連続するテーブル行を収集
    block: list[str] = []
    while i < n and is_table_line(lines[i]):
        block.append(lines[i])
        i += 1

    rows = [split_table_row(b) for b in block]
    if not rows:
        return [], []

    header = rows[0]
    data = []
    for r in rows[1:]:
        if is_separator_row(r):
            continue
        # 列数をヘッダーに合わせて補完/切詰めする
        if len(r) < len(header):
            r = r + [""] * (len(header) - len(r))
        elif len(r) > len(header):
            r = r[: len(header)]
        data.append(r)
    return header, data


def autosize_columns(ws, header: list[str], data: list[list[str]]) -> None:
    for col_idx in range(1, len(header) + 1):
        max_len = len(str(header[col_idx - 1]))
        for row in data:
            cell_val = row[col_idx - 1] if col_idx - 1 < len(row) else ""
            # セル内改行は最長行で評価
            longest = max((len(s) for s in str(cell_val).split("\n")), default=0)
            if longest > max_len:
                max_len = longest
        width = min(MAX_COL_WIDTH, max(MIN_COL_WIDTH, max_len + 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def write_xlsx(
    header: list[str], data: list[list[str]], out_path: Path, sheet_name: str
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31] if sheet_name else "結果出力"

    header_font = Font(bold=True)
    header_fill = PatternFill(
        start_color="FFD9D9D9", end_color="FFD9D9D9", fill_type="solid"
    )
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_align = Alignment(vertical="top", wrap_text=True)

    # ヘッダー
    for col_idx, name in enumerate(header, start=1):
        c = ws.cell(row=1, column=col_idx, value=name)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align

    # データ
    for r_idx, row in enumerate(data, start=2):
        for col_idx in range(1, len(header) + 1):
            val = row[col_idx - 1] if col_idx - 1 < len(row) else ""
            c = ws.cell(row=r_idx, column=col_idx, value=val)
            c.alignment = body_align

    autosize_columns(ws, header, data)
    ws.freeze_panes = "A2"

    if out_path.parent and not out_path.parent.exists():
        out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def convert(in_path: Path, out_path: Path, heading: str, sheet_name: str | None) -> int:
    text = in_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    h_idx = find_heading_line(lines, heading)
    if h_idx >= 0:
        header, data = extract_table(lines, h_idx + 1)
        used_heading = heading
    else:
        sys.stderr.write(
            f"[warn] 見出し「{heading}」が見つかりません。先頭のテーブルを変換します。\n"
        )
        header, data = extract_table(lines, 0)
        used_heading = ""

    if not header:
        sys.stderr.write("[error] 変換対象の Markdown テーブルが見つかりません。\n")
        return 2

    sheet = sheet_name if sheet_name else (used_heading or "結果出力")
    write_xlsx(header, data, out_path, sheet)
    sys.stderr.write(
        f"[done] Excel を出力しました: {out_path} (列数={len(header)}, 行数={len(data)})\n"
    )
    return 0


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="チェック結果レポート(Markdown)の結果テーブルを Excel(.xlsx) へ変換する。"
    )
    p.add_argument(
        "--in",
        dest="in_path",
        required=True,
        help="入力 Markdown レポートのパス（例: specification_check/output/diff_summary_report.md）",
    )
    p.add_argument(
        "--out",
        default="report.xlsx",
        help="出力 Excel パス (既定: report.xlsx)",
    )
    p.add_argument(
        "--heading",
        default=DEFAULT_HEADING,
        help=f"変換対象テーブルの見出しテキスト (既定: {DEFAULT_HEADING})",
    )
    p.add_argument(
        "--sheet",
        default=None,
        help="出力シート名 (既定: 見出しテキスト)",
    )
    return p


def main() -> int:
    args = build_parser().parse_args()
    in_path = Path(args.in_path).expanduser()
    if not in_path.is_file():
        sys.stderr.write(f"[error] 入力 Markdown が見つかりません: {in_path}\n")
        return 2
    return convert(in_path, Path(args.out).expanduser(), args.heading, args.sheet)


if __name__ == "__main__":
    raise SystemExit(main())
